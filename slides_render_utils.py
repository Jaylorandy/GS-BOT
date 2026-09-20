import hashlib
import json
import os
import sys
import tempfile

from PIL import Image, ImageOps
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE


TEMP_IMAGE_VARIANTS = {}
IMAGE_METADATA_CACHE = {}


def ensure_dir(path):
    os.makedirs(path, exist_ok=True)
    return path


def get_render_cache_root():
    candidates = [
        os.path.join(os.path.expanduser('~'), '.gsbot', 'ppt-render-cache'),
        os.path.join(tempfile.gettempdir(), 'gsbot-ppt-render-cache'),
    ]
    for candidate in candidates:
        try:
            ensure_dir(candidate)
            probe_path = os.path.join(candidate, f'.probe-{os.getpid()}')
            with open(probe_path, 'w', encoding='utf-8') as handle:
                handle.write('ok')
            os.remove(probe_path)
            return candidate
        except Exception:
            continue
    return ensure_dir(os.path.join(tempfile.gettempdir(), 'gsbot-ppt-render-cache'))


def build_image_signature(img_path):
    stats = os.stat(img_path)
    return {
        'path': os.path.abspath(img_path),
        'size': int(stats.st_size or 0),
        'mtimeMs': int(round(stats.st_mtime * 1000)),
    }


def build_image_cache_key(img_path):
    payload = json.dumps(build_image_signature(img_path), sort_keys=True, separators=(',', ':'))
    return hashlib.sha1(payload.encode('utf-8')).hexdigest()


def get_image_cache_paths(img_path):
    cache_key = build_image_cache_key(img_path)
    root = get_render_cache_root()
    return {
        'key': cache_key,
        'meta': os.path.join(root, f'{cache_key}.json'),
        'converted': os.path.join(root, f'{cache_key}.png'),
    }


def read_cached_image_metadata(img_path):
    try:
        cache_paths = get_image_cache_paths(img_path)
        meta_path = cache_paths['meta']
        if not os.path.exists(meta_path):
            return None
        with open(meta_path, 'r', encoding='utf-8') as handle:
            payload = json.load(handle)
        if payload.get('signature') != build_image_signature(img_path):
            return None
        converted_path = payload.get('preparedPath') or img_path
        if converted_path != img_path and not os.path.exists(converted_path):
            return None
        return payload
    except Exception:
        return None


def write_cached_image_metadata(img_path, payload):
    try:
        cache_paths = get_image_cache_paths(img_path)
        with open(cache_paths['meta'], 'w', encoding='utf-8') as handle:
            json.dump(payload, handle, ensure_ascii=False, indent=2)
    except Exception:
        return None


def get_image_size(img_path):
    """获取图片原始宽高"""
    metadata = resolve_image_metadata(img_path)
    return metadata.get('width'), metadata.get('height')


def prepare_powerpoint_image(img_path):
    return resolve_image_metadata(img_path).get('preparedPath') or img_path


def resolve_image_metadata(img_path):
    if not img_path or not os.path.exists(img_path):
        return {
            'preparedPath': img_path,
            'width': None,
            'height': None,
        }

    memory_key = build_image_cache_key(img_path)
    cached_memory = IMAGE_METADATA_CACHE.get(memory_key)
    if cached_memory:
        prepared_path = cached_memory.get('preparedPath') or img_path
        if prepared_path == img_path or os.path.exists(prepared_path):
            return cached_memory

    cached_disk = read_cached_image_metadata(img_path)
    if cached_disk:
        IMAGE_METADATA_CACHE[memory_key] = cached_disk
        prepared_path = cached_disk.get('preparedPath')
        if prepared_path and prepared_path != img_path:
            TEMP_IMAGE_VARIANTS[img_path] = prepared_path
        return cached_disk

    ext = os.path.splitext(img_path)[1].lower()
    prepared_path = img_path

    try:
        with Image.open(img_path) as img:
            normalized = ImageOps.exif_transpose(img)
            width, height = normalized.width, normalized.height
            orientation_changed = (width != img.width or height != img.height)
            needs_conversion = ext == '.webp' or orientation_changed
            if needs_conversion:
                cache_paths = get_image_cache_paths(img_path)
                converted_path = cache_paths['converted']
                if not os.path.exists(converted_path):
                    converted = normalized.convert('RGBA') if normalized.mode in ('RGBA', 'LA', 'P') else normalized.convert('RGB')
                    converted.save(converted_path, 'PNG')
                prepared_path = converted_path

        payload = {
            'signature': build_image_signature(img_path),
            'preparedPath': prepared_path,
            'width': width,
            'height': height,
        }
        IMAGE_METADATA_CACHE[memory_key] = payload
        write_cached_image_metadata(img_path, payload)
        if prepared_path != img_path:
            TEMP_IMAGE_VARIANTS[img_path] = prepared_path
        return payload
    except Exception as error:
        if ext == '.webp':
            print(f"  Error converting WEBP image: {error}", file=sys.stderr)
        return {
            'preparedPath': img_path,
            'width': None,
            'height': None,
        }


def add_image_fit(slide, img_path, x, y, max_w, max_h):
    """添加图片，保持原始比例缩放到最大框内（不拉伸）"""
    if not img_path or not os.path.exists(img_path):
        return None

    metadata = resolve_image_metadata(img_path)
    prepared_img_path = metadata.get('preparedPath') or img_path
    orig_w, orig_h = metadata.get('width'), metadata.get('height')
    if not orig_w or not orig_h:
        return None

    ratio = orig_w / orig_h
    target_ratio = max_w / max_h

    if ratio > target_ratio:
        final_w = max_w
        final_h = int(max_w / ratio)
    else:
        final_h = max_h
        final_w = int(max_h * ratio)

    offset_x = x + (max_w - final_w) // 2
    offset_y = y + (max_h - final_h) // 2

    try:
        return slide.shapes.add_picture(prepared_img_path, offset_x, offset_y, final_w, final_h)
    except Exception as error:
        print(f"  Error adding image: {error}", file=sys.stderr)
        return None


def add_text(slide, text, y_key, layout, left_margin, text_width, font_name, color_map, override_text=None):
    """用固定布局添加文本，统一字体（所有段落）"""
    cfg = layout[y_key]
    txt = override_text if override_text is not None else text
    color = color_map.get(cfg['color'])

    box = slide.shapes.add_textbox(
        left_margin,
        Inches(cfg['y']),
        text_width,
        Inches(cfg['h']),
    )
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    tf.clear()
    lines = txt.split('\n')
    for index, line in enumerate(lines):
        para = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        para.text = line
        para.font.name = font_name
        para.font.size = Pt(cfg['size'])
        para.font.bold = cfg['bold']
        para.font.color.rgb = color
        para.space_after = Pt(0)
        para.space_before = Pt(0)
        para.line_spacing = 1.15

    return box


def add_custom_text(slide, text, cfg, font_name, color_map, default_left_margin, default_text_width):
    """按传入配置渲染文本块"""
    if not text:
        return None

    color = color_map.get(cfg.get('color', 'black'))
    x = cfg.get('x', default_left_margin)
    width = cfg.get('w', default_text_width)

    box = slide.shapes.add_textbox(
        x,
        Inches(cfg['y']),
        width,
        Inches(cfg['h']),
    )
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    txt = str(text)

    tf.clear()
    lines = txt.split('\n')
    for index, line in enumerate(lines):
        para = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        para.text = line
        para.font.name = font_name
        para.font.size = Pt(cfg['size'])
        para.font.bold = cfg.get('bold', False)
        para.font.color.rgb = color
        para.space_after = Pt(0)
        para.space_before = Pt(0)
        para.line_spacing = 1.15

    return box
